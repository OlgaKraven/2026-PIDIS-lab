from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"


def main() -> None:
    files = sorted(REPORTS.glob("*.docx"))
    assert len(files) == 20, f"Ожидалось 20 DOCX, найдено {len(files)}"
    results = []
    for path in files:
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        match = re.match(r"(C[34]_S[567]_LR\d{2})_", path.name)
        assert match, path.name
        lab_id = match.group(1).replace("_", "-")
        assert lab_id in text, f"{path.name}: нет ID"
        assert re.search(r"Критерии LMS · \d+ балл(?:а|ов)?", text), f"{path.name}: нет критериев LMS"
        legacy_name = "Moo" + "dle"
        assert legacy_name.lower() not in text.lower(), f"{path.name}: осталось устаревшее название"
        assert len(doc.tables) >= 3, f"{path.name}: недостаточно таблиц"
        assert len(doc.paragraphs) >= 35, f"{path.name}: недостаточно редактируемых абзацев"
        section = doc.sections[0]
        assert 205 < section.page_width.mm < 215 and 292 < section.page_height.mm < 302, f"{path.name}: не A4"
        hyperlinks = 0
        for rel in doc.part.rels.values():
            if rel.reltype.endswith("/hyperlink") and rel.target_ref.startswith("https://"):
                hyperlinks += 1
        assert hyperlinks >= 3, f"{path.name}: недостаточно рабочих гиперссылок"
        numbered = sum(
            1
            for p in doc.paragraphs
            if (p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None)
            or (p.style.element.pPr is not None and p.style.element.pPr.find(qn("w:numPr")) is not None)
        )
        assert numbered >= 4, f"{path.name}: нет настоящей нумерации/маркеров"
        results.append({"id": lab_id, "tables": len(doc.tables), "paragraphs": len(doc.paragraphs), "hyperlinks": hyperlinks, "numbered_paragraphs": numbered})

    output = ROOT / ".qa" / "reports-structure.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"reports": len(results), "status": "ok", "summary": str(output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
