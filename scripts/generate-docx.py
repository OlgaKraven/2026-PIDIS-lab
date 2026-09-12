from pathlib import Path
import json,copy
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
reference=ROOT/'authoring/report-reference.docx'
labs=json.loads((ROOT/'src/data/labs.json').read_text(encoding='utf8'))['labs']
designs=json.loads((ROOT/'authoring/docx-designs.json').read_text(encoding='utf8'))
out=ROOT/'public/reports';out.mkdir(exist_ok=True)
def font(p,bold=False):
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=bold if bold else r.bold
def para(doc,text='',style=None):
 p=doc.add_paragraph(text,style);font(p);p.paragraph_format.space_after=Pt(6);return p
def heading(doc,text):
 p=para(doc,text,'Heading 1');font(p,True);p.paragraph_format.keep_with_next=True
def addtable(doc,cols,n=3):
 t=doc.add_table(rows=1,cols=len(cols));t.autofit=False
 for c in t.columns:c.width=Cm(16.5/len(cols))
 tr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');tr.append(repeat)
 for i,c in enumerate(t.rows[0].cells):
  c.text=cols[i];font(c.paragraphs[0],True);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EEF7');c._tc.get_or_add_tcPr().append(shade)
 for j in range(n):
  for c in t.add_row().cells:
   c.text='\n';font(c.paragraphs[0])
 for row in t.rows:
  for cell in row.cells:
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D5DDE5');borders.append(edge)
   cell._tc.get_or_add_tcPr().append(borders)
 return t
for lab,d in zip(labs,designs):
 doc=Document(reference);body=doc.element.body
 for e in list(body)[19:]:
  if e.tag!=qn('w:sectPr'):body.remove(e)
 replacements={'Предложение улучшения':lab['title'],'Отчёт по лабораторной работе № 3':'Отчёт по лабораторной работе № '+str(lab['number']),'[Курс] · 2 семестр · ЛР03':f"{3 if lab['semester']<7 else 4} курс · {lab['semester']} семестр · ЛР{lab['number']:02}",'[Название дисциплины]':'МДК.05.01 Проектирование и дизайн информационных систем','[Код и название специальности]':'09.02.07 Информационные системы и программирование','[Город], [Год]':'[Город], 2026'}
 for e in body.iter(qn('w:t')):
  if e.text:
   for a,b in replacements.items():e.text=e.text.replace(a,b)
 for style in doc.styles:
  if style.type in (1,2):
   style.font.name='Times New Roman';style.font.size=Pt(12);style.font.color.rgb=RGBColor(0,0,0)
 for e in list(body.iter(qn('w:fitText'))):e.getparent().remove(e)
 for p in doc.paragraphs:font(p)
 for t in doc.tables:
  for row in t.rows:
   for c in row.cells:
    for p in c.paragraphs:font(p)
 heading(doc,lab['id']+' '+lab['title'])
 para(doc,f"Семестр {lab['semester']} · Максимум {lab['points']} баллов")
 para(doc,'ФИО ____________________ Группа __________ Вариант ________')
 para(doc,lab['goal'])
 para(doc,'Используйте CSV из папки Данные. В ответах приводите ID записей; добавляйте строки по необходимости. Версия набора: '+lab['version']+'.')
 for title,cols in d['sections'][:2]:
  heading(doc,title);addtable(doc,cols,2)
 doc.add_page_break()
 if len(d['sections'])>2:
  title,cols=d['sections'][2];heading(doc,title);addtable(doc,cols,2)
 else:
  heading(doc,'Проверка решения');para(doc,d['steps'][-1][2]);addtable(doc,['Случай и ID','Проверка','Результат и изменение'],2)
 heading(doc,'Обоснованный вывод')
 para(doc,lab['professionalChoice']);para(doc,'Запишите решение, подтверждающий факт и условие, при котором решение потребуется пересмотреть.');para(doc,'\n')
 heading(doc,'Критерии оценивания')
 t=addtable(doc,['Результат','Что оценивается','Баллы'],0)
 for row in lab['rubric']:
  cells=t.add_row().cells
  for c,value in zip(cells,row):c.text=str(value);font(c.paragraphs[0])
 for rel in list(doc.part.rels.values()):
  if rel.is_external:doc.part.drop_rel(rel.rId)
 doc.save(out/lab['reportFile'])
print('Created',len(labs),'DOCX forms')
